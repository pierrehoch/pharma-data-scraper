"""Full public Chinese catalogues and lossless, separately audited price extraction."""

from __future__ import annotations

import hashlib
import csv
import json
import re
import shutil
import subprocess
import unicodedata
import zipfile
import zlib
from collections import Counter
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import NamedTuple
from urllib.parse import unquote, urlencode, urljoin, urlparse
from xml.etree import ElementTree

import openpyxl
from bs4 import BeautifulSoup
from openpyxl.utils.exceptions import InvalidFileException
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfminer.pdfparser import PDFSyntaxError
from pdfplumber.utils.exceptions import PdfminerException
from xlrd.biffh import XLRDError

from app.adapters._cn_provincial import COLUMNS, UNITS, chinese_pack
from app.core.enums import ErrorType, PriceBasis, PriceStatus, ProductLevel
from app.core.errors import DiscoveryError, DownloadError, ParserError, SchemaChangedError
from app.core.parsing import clean_str, profile_rows
from app.core.quality import NOT_CURRENT_FLAG
from app.core.snapshots import DownloadedFile
from app.core.types import ArtifactFile, NormalizedRecord, ParsedDataset, ParsedRecord, PriceFact, RawSnapshot, SourceArtifact, SourceProductData
from app.observability import Progress, get_logger

log = get_logger(__name__)
GD = "https://www.gdmede.com.cn"
HB = "https://ybj.hubei.gov.cn"
HB_CATALOGUE = HB + "/zfxxgk/zc/qtzdgkwj/zcwj.json"
GD_QUERY = GD + "/api/service/nationalCentralizedProcurement/ec/medicines"
DRUG_RE = re.compile(r"药品|中成药|中药|药价|药物|集采|集中采购")
NON_DRUG_RE = re.compile(r"耗材|医疗器械|医疗设备|医疗服务|医药机构|消毒产品|营养|诊断试剂|检测试剂")
PRICE_RE = re.compile(r"价格|挂网价|中选价|中标价|成交价|出厂价|结算价|协议价|医保价|投标价|批发价|销售价|报价|采购价|零售价|支付标准|限价|参考价|单价|单片价|单盒价|入市价|最低价|最高价|中位价|日均费用|日费用|金额")
NON_PRICE_RE = re.compile(
    r"方式|类型|说明|依据|状态|日期|政策|是否|备注|标识|标志|变化率|降幅|系数|比例|比值|比率|占比|权重|"
    r"排名|排序|名次|得分|分数|评分|分值|价格单位|报价单位|折扣率|"
    r"(?:企业|公司|厂商|主体|名称|代码|编码|编号|数量)$"
)
DATE_RE = r"(\d{4})[年./-](\d{1,2})[月./-](\d{1,2})日?"
ALIASES = {
    "code": ["国家医保代码", "国家医保编码", "医保编码", "药品代码", "药品编码", "医保药品编码",
             "医保统一编码", "药品统一编码", "医保代码", "医保药品分类与代码", "国家药品代码", "价格治理医保编码"],
    "local_code": ["药品ID", "药品id", "流水号", "产品编号", "交易编码", "药品编号"],
    "name": ["药品通用名", "药品通用名称", "药品名称", "通用名", "通用名称", "产品名称", "产品名", "药品名", "品名", "品种名称", "品种名", "药物名称"],
    "form": ["剂型名称", "剂型", "合并归类剂型"],
    "strength": ["具体规格", "规格", "药品规格"],
    "pack": ["包装规格", "规格包装", "包装"],
    "quantity": ["最小包装数量", "包装数量", "包装转换比", "转换比"],
    "unit": ["最小制剂单位", "制剂单位", "最小使用单位", "最小计量单位", "最小制剂单位（片/粒/袋/支等）"],
    "pricing_unit": ["计价单位", "计量单位", "单位", "包装计量单位", "最小计量单位（价格单位）"],
    "pack_unit": ["最小包装单位", "包装单位"],
    "manufacturer": ["生产企业/厂商", "生产企业名称", "生产企业", "生产厂家", "生产厂商"],
    "authorisation": ["批准文号/注册证号", "批准文号", "批准文号（注册证号）", "注册证号"],
}
TABLE_ERRORS = (ParserError, SchemaChangedError, ValueError, OSError, KeyError, zipfile.BadZipFile,
                InvalidFileException, XLRDError, PDFSyntaxError, PDFPasswordIncorrect, PdfminerException,
                ElementTree.ParseError, zlib.error)


def compact(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("(", "（").replace(")", "）")


def is_price_column(label: str) -> bool:
    return bool(PRICE_RE.search(label)) and not bool(NON_PRICE_RE.search(label))


def non_drug_table(header: list[str], meta: dict) -> bool:
    device_headers = {"体外诊断试剂代码", "医用耗材代码", "医保耗材代码", "医用耗材编码",
                      "医疗器械注册证号", "耗材分类与代码（20位）"}
    label = meta.get("label", "")
    title = meta.get("title", "")
    return bool(device_headers.intersection(header)) or bool(
        re.search(r"耗材|医疗器械|医疗设备|医疗服务|诊断试剂|检测试剂", label)
        and not re.search(r"药品|中成药|中药|药物|药价", label)
    ) or bool(NON_DRUG_RE.search(title) and "药" not in title)


def original_amount(value: object) -> tuple[str | None, str | None]:
    """Do not turn explanatory text, a range, or a coded category into a number."""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        amount = Decimal(str(value))
        if not amount.is_finite():
            return None, f"Nonfinite price cell: {value!r}"
        return str(amount), None if amount > 0 else "Nonpositive disclosed amount; not normalized"
    text = unicodedata.normalize("NFKC", str(value)).strip()
    match = re.fullmatch(
        r"(?P<qualifier>[≤≥<>≈约]?)\s*[¥￥]?\s*(?P<amount>[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)"
        r"\s*(?P<unit>元(?:/(?:片|粒|丸|袋|支|瓶|盒|贴|包|g|ml|mg))?|\(元\))?",
        text,
    )
    if not match:
        return None, f"Non-scalar price cell: {value!r}"
    amount = Decimal(match["amount"].replace(",", ""))
    if amount <= 0:
        return str(amount), "Nonpositive disclosed amount; not normalized"
    if match["qualifier"] or "/" in (match["unit"] or ""):
        return str(amount), f"Qualified original amount, not automatically normalized: {text}"
    return str(amount), None


def drug_notice(title: str) -> bool:
    return bool(DRUG_RE.search(title)) and (bool(re.search(r"药品|中成药|中药|药价", title)) or not bool(NON_DRUG_RE.search(title)))


def official(url: str, host: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or parsed.netloc != host:
        raise DiscoveryError(f"Non-official public URL: {url}")
    return parsed._replace(scheme="https", fragment="").geturl()


def notice_metadata(title: str, text: str, published: str | None) -> dict:
    proposed = bool(re.search(r"拟|公示|征求意见", title))
    withdrawn = bool(re.search(r"废止|作废|撤网|取消.*(?:中选|挂网|资格)|停止|暂停.*(?:挂网|交易)", title))
    dates = {date(int(y), int(m), int(d)).isoformat() for y, m, d in re.findall(
        rf"自?{DATE_RE}起?(?:执行|实施|生效)", text,
    )}
    return {
        "title": title, "published_at": published, "status": "proposed" if proposed else "withdrawn" if withdrawn else "final",
        "effective_from": next(iter(dates)) if len(dates) == 1 and not proposed else None,
    }


class PriceTable(NamedTuple):
    name: str
    header: list[str]
    rows: list
    caption: str = ""


def price_semantics(label: str, unit: str | None, caption: str, dosage_unit: str | None = None) -> dict:
    currency = "CNY" if "元" in label or re.search(r"(?:金额单位|价格单位|币种|单位)[：:]?(?:人民币)?元|人民币元", compact(caption)) else None
    explicit_unit = re.search(r"元[/／]([片粒丸袋支瓶盒贴包])(?=[）)]|$)", label)
    minimum_unit = bool(re.search(r"最小制剂单位|最小使用单位|最小计量单位", label))
    if explicit_unit:
        unit = explicit_unit[1]
    elif minimum_unit:
        unit = dosage_unit
    if minimum_unit or "包装计量单位" in label:
        fact_unit = UNITS.get(unit)
    else:
        fact_unit = "pack" if "包装" in label or unit in {"盒", "箱", "包"} else UNITS.get(unit)
    aggregate = "金额" in label and "单价" not in label
    daily_cost = bool(re.search(r"日均费用|日费用", label))
    benchmark = bool(re.search(r"基准|最低|参考|中位|同类品种", label))
    if aggregate or daily_cost:
        basis, variant, fact_unit = None, None, None
    elif "零售" in label:
        basis, variant = PriceBasis.RETAIL, "retail_published_price"
    elif re.search(r"支付标准|报销", label):
        basis, variant = PriceBasis.REIMBURSEMENT, "reimbursement_published_standard"
    elif "挂网" in label:
        basis, variant = PriceBasis.PUBLIC_PROCUREMENT, "guawang_pack_listing" if fact_unit == "pack" else "guawang_unit_listing"
    elif re.search(r"中选|中标", label):
        basis = PriceBasis.PUBLIC_PROCUREMENT
        if benchmark:
            variant = "tender_price_benchmark"
        elif re.search(r"入围|备选", label):
            variant = "selected_or_shortlisted_published_price"
        else:
            variant = "awarded_pack_price" if fact_unit == "pack" else "awarded_unit_price"
    elif "采购" in label:
        basis, variant = PriceBasis.PUBLIC_PROCUREMENT, "procurement_published_price"
    else:
        basis, variant = None, None
    return {"currency": currency, "pricing_unit": unit, "fact_unit": fact_unit, "price_type": basis.value if basis else None,
            "price_variant": variant, "price_status": "ceiling" if re.search(r"最高|限价", label) and "同类品种" not in label else "reference",
            "value_kind": "aggregate_amount" if aggregate else "daily_cost" if daily_cost else "benchmark_price" if benchmark else "price"}


NOT_CURRENT_REASON = "historical/proposed/withdrawn/future value"
UNESTABLISHED_REASON = "currency/price unit or concept not explicitly established"
QUERY_REASON = "Public query does not establish currency, price unit or effective date"
_NOT_CURRENT_STATUS_FLAGS = {"previous": "previous_value", "proposed": "proposed_value", "withdrawn": "withdrawn_value"}


def publication_flags(data: dict) -> list[str] | None:
    """Warnings for publishing a positive source price; None when the value must stay out of the database.

    Established values get no flag. Values whose only gaps are currency/basis/concept/date or a
    non-current status are published with explicit flags; nothing missing is guessed except CNY.
    Templates, raw herbs, aggregates, daily costs and malformed or nonpositive cells stay unpublished.
    """
    if data.get("normalized"):
        return []
    reason = data.get("reason")
    if reason not in {NOT_CURRENT_REASON, UNESTABLISHED_REASON, QUERY_REASON}:
        return None
    if data.get("value_kind") not in {"price", "benchmark_price"} or data.get("amount") is None:
        return None
    if Decimal(str(data["amount"])) <= 0:
        return None
    flags = []
    if data.get("currency") != "CNY":
        flags.append("currency_inferred_cny")
    if not data.get("fact_unit"):
        flags.append("price_basis_unknown")
    if not data.get("price_type"):
        flags.append("price_concept_unknown")
    if not data.get("published_at") and not data.get("effective_from"):
        flags.append("publication_date_unknown")
    if reason == NOT_CURRENT_REASON:
        flags += [_NOT_CURRENT_STATUS_FLAGS.get(data.get("status"), "future_value"), NOT_CURRENT_FLAG]
    return flags


class PublicChinaMixin:
    """Uses the existing SourceAdapter HTTP/store contract; explicit documents retain reviewed mode."""

    parser_type = "public JSON + document tables"
    parser_version = "5"

    def discover(self) -> SourceArtifact:
        replay = self.option("public_snapshot_dir")
        if replay is not None:
            if self.check_configuration():
                from app.core.errors import SourceUnsupportedError

                raise SourceUnsupportedError(self.check_configuration()[0].message)
            if self.config.coverage_scope != f"province:{self.province}":
                raise DiscoveryError("Chinese public snapshot coverage must remain provincial")
            snapshot = self.ctx.store.load(Path(replay))
            if snapshot.source_id != self.source_id or not any(f.role == "public_manifest" for f in snapshot.files):
                raise DiscoveryError("public_snapshot_dir must be a verified public snapshot of this source")
            return SourceArtifact(self.source_id, [
                ArtifactFile((snapshot.directory / "metadata.json").resolve().as_uri(), "metadata.json", role="public_replay"),
            ], discovered_from=str(snapshot.directory), notes=["Verified local public snapshot replay; no network discovery"],
                partial=snapshot.partial)
        if self.option("documents") is not None:
            return super().discover()
        return self.discover_public()

    def download(self, artifact: SourceArtifact) -> RawSnapshot:
        replay = next((f for f in artifact.files if f.role == "public_replay"), None)
        if replay is not None:
            return self.ctx.store.load(Path(unquote(urlparse(replay.url).path)).parent)
        if any(f.role == "public_manifest" for f in artifact.files):
            return self.download_public(artifact)
        return super().download(artifact)

    def parse(self, snapshot: RawSnapshot) -> ParsedDataset:
        if any(f.role == "public_manifest" for f in snapshot.files):
            return self.parse_public(snapshot)
        return super().parse(snapshot)

    def normalize_record(self, record: ParsedRecord, dataset: ParsedDataset) -> NormalizedRecord:
        if "extraction_directory" in dataset.context:
            return self.normalize_public(record)
        return super().normalize_record(record, dataset)

    def _local_artifact(self, folder: Path, name: str, value: object, role: str) -> ArtifactFile:
        path = folder / name
        path.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True), encoding="utf-8")
        return ArtifactFile(path.resolve().as_uri(), name, role=role)

    def _catalogue_json(self, url: str) -> object:
        return self.ctx.http.get_json(url)

    def _capture(self, folder: Path, value: object, role: str) -> ArtifactFile:
        return self._local_artifact(folder, f"{role}.json", value, role)

    def _gd_notices(self, folder: Path, files: list[ArtifactFile], report: dict) -> list[dict]:
        notices = []
        seen: set[str] = set()
        total = None
        page = 1
        while True:
            url = GD + "/api/announcement/paginateAnnouncement?" + urlencode({"page": page, "count": 100, "label": "collect"})
            response = self._catalogue_json(url)
            data = response.get("data") if isinstance(response, dict) else None
            if not isinstance(data, dict) or not isinstance(data.get("array"), list) or not isinstance(data.get("total"), int):
                raise DiscoveryError(f"Guangdong invalid announcement pagination at page {page}")
            if data.get("page") != page or data["total"] < 0:
                raise DiscoveryError("Guangdong catalogue returned the wrong page or total")
            if total is None:
                total = data["total"]
            if data["total"] != total:
                raise DiscoveryError("Guangdong catalogue total changed during collection; retry a consistent catalogue")
            files.append(self._capture(folder, response, f"catalogue_{page:04d}"))
            items = data["array"]
            if not items and len(seen) < total:
                raise DiscoveryError("Guangdong prematurely empty catalogue page")
            for item in items:
                if not isinstance(item, dict) or not isinstance(item.get("id"), str) or not isinstance(item.get("title"), str):
                    raise DiscoveryError("Guangdong invalid announcement identity")
                if item["id"] in seen:
                    raise DiscoveryError("Guangdong repeated announcement across pages")
                seen.add(item["id"])
                notices.append(item)
            if len(seen) >= total:
                break
            page += 1
        if len(seen) != total:
            raise DiscoveryError(f"Guangdong expected {total} notices, received {len(seen)}")
        report["catalogue_total"] = total
        return notices

    def _gd_query(self, folder: Path, files: list[ArtifactFile], report: dict) -> None:
        total = None
        seen: set[str] = set()
        page = 1
        while True:
            url = GD_QUERY + "?" + urlencode({"rows": 50, "currentPage": page, "searchKey": ""})
            data = self._catalogue_json(url)
            if not isinstance(data, dict) or not isinstance(data.get("list"), list):
                raise DiscoveryError("Invalid Guangdong public procurement response")
            if data.get("currentPage") != page or not isinstance(data.get("totalCount"), int):
                raise DiscoveryError("Guangdong procurement returned an invalid page/total")
            if total is None:
                total = data["totalCount"]
            if total != data["totalCount"] or total < 0:
                raise DiscoveryError("Guangdong procurement total changed; refusing incomplete extraction")
            rows = data["list"]
            if not rows and len(seen) < total:
                raise DiscoveryError("Guangdong procurement prematurely empty page")
            for row in rows:
                if not isinstance(row, dict) or not row.get("id") or str(row["id"]) in seen:
                    raise DiscoveryError("Guangdong procurement invalid or repeated product identity")
                seen.add(str(row["id"]))
            files.append(self._capture(folder, data, f"query_{page:04d}"))
            if len(seen) >= total:
                break
            page += 1
        if len(seen) != total:
            raise DiscoveryError(f"Guangdong procurement expected {total} records, received {len(seen)}")
        report["public_query_rows"] = len(seen)

    def _notice(self, item: dict, folder: Path, files: list[ArtifactFile], index: int) -> tuple[dict, list[dict]]:
        if self.province == "CN-GD":
            url = GD + "/announcement/announcement/detail?id=" + item["id"]
            detail = getattr(self, "_gd_cache", {}).get(item["id"])
            if detail is None:
                detail = self._catalogue_json(GD + "/api/announcement/detailAnnouncement?" + urlencode({"id": item["id"]}))
            data = detail.get("data") if isinstance(detail, dict) else None
            if not isinstance(data, dict) or not isinstance(data.get("content"), str):
                raise DiscoveryError(f"Invalid Guangdong notice: {url}")
            files.append(self._capture(folder, detail, f"notice_{index:05d}"))
            text = BeautifulSoup(data["content"], "html.parser").get_text(" ", strip=True)
            meta = notice_metadata(str(data.get("title") or item["title"]), text, data.get("publishTime"))
            attachments = list(data.get("attachments") or [])
            soup = BeautifulSoup(data["content"], "html.parser")
            attachments += [{"url": a["href"], "fileName": a.get_text(" ", strip=True)} for a in soup.select("a[href]")]
        else:
            url = official(item["URL"], self.host)
            archive = getattr(self, "_browser_notices", {})
            captured = archive.get(url)
            if captured is not None:
                if captured.get("status") != 200:
                    raise DownloadError(f"Browser-captured public notice unavailable: {url}: {captured.get('error')}")
                html = captured["html"]
            else:
                html = self.ctx.http.get_text(url).text
            files.append(self._local_artifact(folder, f"notice_{index:05d}.json",
                                              {"url": url, "html": html}, f"notice_{index:05d}"))
            soup = BeautifulSoup(html, "html.parser")
            text = soup.get_text(" ", strip=True)
            meta = notice_metadata(item["FILENAME"], text, (item.get("DOCRELTIME") or "")[:10] or None)
            attachments = (captured.get("attachments") or []) if captured is not None else []
            attachments += [{"url": a["href"], "fileName": a.get_text(" ", strip=True)} for a in soup.select("a[href]")]
        meta["notice_url"] = url
        meta["text"] = text
        return meta, attachments

    def discover_public(self) -> SourceArtifact:
        if self.check_configuration():
            from app.core.errors import SourceUnsupportedError

            raise SourceUnsupportedError(self.check_configuration()[0].message)
        if self.config.coverage_scope != f"province:{self.province}":
            raise DiscoveryError("Chinese public catalogue coverage must remain provincial")
        folder = self.ctx.store.staging_dir(self.source_id)
        files: list[ArtifactFile] = []
        report = {"scope": "all public drug notices in source catalogue; no authentication bypass",
                  "failures": [], "attachments": [], "notices_without_attachments": [], "excluded_notices": []}
        if self.province == "CN-GD":
            cache_path = self.option("notice_cache_dir")
            if cache_path is not None:
                self._gd_cache = {}
                for path in Path(cache_path).glob("notice_*.json"):
                    saved = json.loads(path.read_text(encoding="utf-8"))
                    self._gd_cache[str(saved["data"]["id"])] = saved
                report["cached_notices_available"] = len(self._gd_cache)
            notices = self._gd_notices(folder, files, report)
            self._gd_query(folder, files, report)
        else:
            archive_path = self.option("browser_notice_archive")
            if archive_path is not None:
                archive_path = Path(archive_path)
                archive = json.loads(archive_path.read_text(encoding="utf-8"))
                if not isinstance(archive, dict) or not isinstance(archive.get("notices"), list):
                    raise DiscoveryError("browser_notice_archive must contain a notices list exported from public pages")
                self._browser_notices = {official(n["url"], self.host): n for n in archive["notices"]}
                files.append(self._capture(folder, archive, "browser_notice_archive"))
            data = self._catalogue_json(HB_CATALOGUE)
            if not isinstance(data, dict) or not isinstance(data.get("data"), list) or not data["data"]:
                raise DiscoveryError("Invalid/empty Hubei complete public catalogue")
            files.append(self._capture(folder, data, "catalogue_0001"))
            notices = data["data"]
            if any(not isinstance(n, dict) or not isinstance(n.get("FILENAME"), str) or not isinstance(n.get("URL"), str) for n in notices):
                raise DiscoveryError("Hubei catalogue contains invalid notice identities")
            urls = [urlparse(n["URL"])._replace(scheme="https", fragment="").geturl() for n in notices]
            if len(set(urls)) != len(urls):
                raise DiscoveryError("Hubei public catalogue repeats notice URLs")
            report["catalogue_total"] = len(notices)
        report["notices_examined"] = len(notices)
        seen: set[str] = set()
        candidates = []
        for index, item in enumerate(notices):
            title = item["title"] if self.province == "CN-GD" else item["FILENAME"]
            if self.province != "CN-GD" and urlparse(item["URL"]).hostname != self.host:
                report["excluded_notices"].append({"title": title, "url": item["URL"], "reason": "outside official source host"})
                continue
            candidates.append((index, item))
        progress = Progress(f"{self.source_id} notices", total=len(candidates), unit="notices", check_every=1)

        def fetch_notice(candidate):
            index, item = candidate
            captured = []
            try:
                meta, attachments = self._notice(item, folder, captured, index)
            except (DownloadError, DiscoveryError) as exc:
                return item, captured, None, [], exc
            return item, captured, meta, attachments, None

        with ThreadPoolExecutor(max_workers=4) as executor:
            results = list(executor.map(fetch_notice, candidates))
        for item, captured, meta, attachments, error in results:
            title = item["title"] if self.province == "CN-GD" else item["FILENAME"]
            files.extend(captured)
            progress.update(attachments=len(seen), failures=len(report["failures"]))
            if error is not None:
                log.warning("%s notice unavailable: %s: %s", self.source_id, title, error)
                report["failures"].append({"title": title, "url": item.get("URL") or GD + "/announcement/announcement/detail?id=" + item["id"], "reason": str(error)})
                continue
            found = False
            for attachment in attachments:
                href = attachment.get("url")
                if not isinstance(href, str):
                    continue
                path = urlparse(urljoin(meta["notice_url"], href)).path.lower()
                suffix = Path(path).suffix
                if suffix not in {".xlsx", ".xls", ".pdf", ".csv", ".zip", ".rar", ".7z", ".doc", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
                    continue
                try:
                    url = official(urljoin(meta["notice_url"], href), self.host)
                except DiscoveryError as exc:
                    report["failures"].append({"url": href, "reason": str(exc)})
                    continue
                # filename query parameters do not identify different attachment bytes.
                identity = urlparse(url)._replace(query="").geturl()
                if identity in seen:
                    continue
                seen.add(identity)
                found = True
                role = f"attachment_{len(report['attachments']):05d}"
                label = str(attachment.get("fileName") or Path(path).name)
                entry = {**meta, "url": url, "role": role, "label": label, "suffix": suffix}
                if non_drug_table([], entry):
                    entry["excluded_reason"] = "non-drug attachment"
                report["attachments"].append(entry)
                if not entry.get("excluded_reason"):
                    files.append(ArtifactFile(url, f"{role}{suffix}", role=role))
            if not found:
                report["notices_without_attachments"].append({"url": meta["notice_url"], "title": title})
        progress.finish(attachments=len(seen), failures=len(report["failures"]))
        files.append(self._capture(folder, report, "public_manifest"))
        return SourceArtifact(self.source_id, files, discovered_from=self.config.source_url,
                              notes=["Full public catalogue enumeration; coverage/exclusions recorded in public_manifest.json.",
                                     "Internal research, not a licence for commercial reuse."],
                              partial=bool(report["failures"]))

    def download_public(self, artifact: SourceArtifact) -> RawSnapshot:
        folder = self.ctx.store.staging_dir(self.source_id)
        downloaded = []
        failures = []
        def fetch_file(indexed):
            index, file = indexed
            try:
                result = self.ctx.http.download(file.url, folder / f"{index:05d}_{file.filename}")
            except DownloadError as exc:
                if not file.role.startswith("attachment_"):
                    raise
                log.warning("Attachment download excluded: %s: %s", file.url, exc)
                return None, {"url": file.url, "role": file.role, "reason": str(exc)}
            return DownloadedFile(file, result.path, result.final_url, result.content_type, file.filename), None

        with ThreadPoolExecutor(max_workers=4) as executor:
            for result, failure in executor.map(fetch_file, enumerate(artifact.files)):
                if failure is not None:
                    failures.append(failure)
                else:
                    downloaded.append(result)
        meta = self._capture(folder, {"failures": failures}, "download_report")
        downloaded.append(DownloadedFile(meta, Path(unquote(urlparse(meta.url).path)), meta.url, "application/json"))
        artifact.partial |= bool(failures)
        return self.ctx.store.save(self.source_id, artifact, downloaded, parser_version=self.parser_version,
                                   previous_id=self.ctx.previous_snapshot_id, previous_hashes=self.ctx.previous_snapshot_hashes)

    def parse_public(self, snapshot: RawSnapshot) -> ParsedDataset:
        manifest = json.loads(snapshot.file("public_manifest").path.read_text(encoding="utf-8"))
        downloads = json.loads(snapshot.file("download_report").path.read_text(encoding="utf-8"))
        records: list[ParsedRecord] = []
        extracted = []
        exclusions = list(manifest["failures"]) + downloads["failures"]
        counts = []
        for file in snapshot.files_with_role_prefix("notice_"):
            saved = json.loads(file.path.read_text(encoding="utf-8"))
            if self.province == "CN-GD":
                notice = saved["data"]
                html = notice["content"]
                url = GD + "/announcement/announcement/detail?id=" + notice["id"]
                meta = notice_metadata(notice["title"], BeautifulSoup(html, "html.parser").get_text(" ", strip=True),
                                       notice.get("publishTime"))
            else:
                html = saved["html"]
                url = saved["url"]
                meta = next((m for m in manifest["attachments"] if m["notice_url"] == url), None)
                if meta is None:
                    meta = notice_metadata("", BeautifulSoup(html, "html.parser").get_text(" ", strip=True), None)
            meta = {**meta, "url": url, "notice_url": url}
            for index, table in enumerate(BeautifulSoup(html, "html.parser").select("table")):
                rows = [[cell.get_text(" ", strip=True) for cell in row.select("td, th")] for row in table.select("tr")]
                for header_index, header in enumerate(rows[:10]):
                    compact_header = [compact(v) for v in header]
                    if not any(is_price_column(t) for t in compact_header):
                        continue
                    if non_drug_table(compact_header, meta):
                        exclusions.append({"url": url, "table": index, "reason": "Non-drug medical-device/service table"})
                        break
                    new_records, values = extract_table(compact_header, list(enumerate(rows[header_index + 1:], header_index + 2)),
                                                       meta, f"{file.filename}:table{index}")
                    if values:
                        records.extend(new_records)
                        extracted.extend(values)
                        break
        for file in snapshot.files_with_role_prefix("query_"):
            data = json.loads(file.path.read_text(encoding="utf-8"))
            for row in data["list"]:
                amount, amount_issue = original_amount(row.get("bidding_price"))
                ref = f"{file.filename}:id={row['id']}"
                raw = {"amount": amount, "original_value": row.get("bidding_price"), "value_kind": "price",
                       "price_label": "bidding_price", "pricing_unit": None,
                       "currency": None, "status": "published_status_unspecified",
                       "published_at": None, "effective_from": None,
                       "source_url": GD_QUERY + "?" + urlencode({"rows": 50, "currentPage": data["currentPage"], "searchKey": ""}),
                       "notice_url": GD + "/supportAndService/nationalCentralizedProcurement",
                       "row_ref": ref, "raw": row, "normalized": False,
                       "reason": amount_issue or QUERY_REASON}
                raw["product"] = {
                    "code": f"GD_QUERY:{row['id']}", "name": row.get("genericName"), "form": row.get("dosageForms"),
                    "strength": row.get("strengths"), "pack": row.get("packing"), "manufacturer": row.get("manufacturer"),
                    "authorisation": row.get("approval_number"),
                }
                extracted.append(raw)
                records.append(ParsedRecord(ref, f"GD_QUERY:{row['id']}", {
                    **raw["product"], "amount": amount, "raw": raw,
                }, role="presentations"))
        by_role = {f.role: f for f in snapshot.files}
        progress = Progress(f"{self.source_id} price documents", total=len(manifest["attachments"]),
                            unit="documents", check_every=1)
        for meta in manifest["attachments"]:
            progress.update(price_cells=len(extracted))
            if meta.get("excluded_reason"):
                exclusions.append({"url": meta["url"], "reason": meta["excluded_reason"]})
                continue
            file = by_role.get(meta["role"])
            if file is None:
                continue
            if non_drug_table([], meta):
                exclusions.append({"url": meta["url"], "reason": "Non-drug medical-device/service attachment"})
                continue
            local_records, local_extracted = [], []
            member_exclusions = []
            try:
                for table in price_tables(file.path, meta["suffix"], exclusions=member_exclusions):
                    table_meta = {**meta, "table_caption": table.caption}
                    if non_drug_table(table.header, table_meta):
                        exclusions.append({"url": meta["url"], "table": table.name, "reason": "Non-drug medical-device/service table"})
                        continue
                    new_records, prices = extract_table(table.header, table.rows, table_meta, f"{file.filename}:{table.name}")
                    local_records.extend(new_records)
                    local_extracted.extend(prices)
                if not local_extracted:
                    exclusions.append({"url": meta["url"], "reason": "No supported numeric drug-price table"})
            except TABLE_ERRORS as exc:
                log.warning("Price document excluded: %s: %s", meta["url"], exc)
                exclusions.append({"url": meta["url"], "reason": str(exc)})
            exclusions.extend({"url": meta["url"], **e} for e in member_exclusions)
            records.extend(local_records)
            extracted.extend(local_extracted)
            counts.append({"url": meta["url"], "price_cells": sum(p["amount"] is not None for p in local_extracted),
                           "unparsed_cells": sum(p["amount"] is None for p in local_extracted)})
        progress.finish(price_cells=len(extracted))
        for price in extracted:
            flags = publication_flags(price)
            price["published"] = flags is not None
            price["quality_flags"] = ";".join(flags or [])
        # Derived results are separate from immutable raw snapshots; every amount retains its original column.
        output = self.ctx.settings.raw_storage_dir.parent / "extractions" / self.source_id / snapshot.directory.name
        output.mkdir(parents=True, exist_ok=True)
        with (output / "prices.jsonl").open("w", encoding="utf-8") as handle:
            for price in extracted:
                handle.write(json.dumps(price, ensure_ascii=False, default=str) + "\n")
        product_columns = ["code", "name", "form", "strength", "pack", "manufacturer", "authorisation"]
        price_columns = ["amount", "original_value", "value_kind", "currency", "pricing_unit", "price_type",
                         "price_status", "price_variant", "fact_unit", "price_label", "status", "published_at",
                         "effective_from", "normalized", "reason", "published", "quality_flags",
                         "source_url", "notice_url", "row_ref"]
        with (output / "prices.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["source_id", *product_columns, *price_columns])
            writer.writeheader()
            for price in extracted:
                writer.writerow({"source_id": self.source_id, **price["product"],
                                 **{key: price.get(key) for key in price_columns}})
        summary = {
            "snapshot_id": snapshot.id, "extraction_parser_version": self.parser_version,
            "catalogue_total": manifest["catalogue_total"],
            "notices_examined": manifest["notices_examined"], "attachments_discovered": len(manifest["attachments"]),
            "public_query_rows": manifest.get("public_query_rows", 0),
            "extracted_price_cells": sum(r["amount"] is not None for r in extracted),
            "unparsed_price_cells": sum(r["amount"] is None for r in extracted),
            "normalizable_price_cells": sum(bool(r.get("normalized")) for r in extracted),
            "published_price_cells": sum(r["published"] for r in extracted),
            "published_with_warnings_cells": sum(r["published"] and not r.get("normalized") for r in extracted),
            "published_quality_flags": dict(Counter(f for r in extracted if r["quality_flags"]
                                                    for f in r["quality_flags"].split(";"))),
            "positive_price_cells": sum(r["amount"] is not None and Decimal(r["amount"]) > 0 for r in extracted),
            "price_cell_statuses": dict(Counter(r["status"] for r in extracted if r["amount"] is not None)),
            "numeric_value_kinds": dict(Counter(r["value_kind"] for r in extracted if r["amount"] is not None)),
            "notice_details_captured": len(snapshot.files_with_role_prefix("notice_")),
            "distinct_source_product_codes": len({r.key for r in records}),
            "documents": counts, "exclusions": exclusions,
            "scope": "Public drug prices, historical/proposed values kept in extraction; not a current national tariff",
        }
        scan_groups = {}
        for exclusion in exclusions:
            if exclusion.get("reason") != "Image-only PDF page; requires reviewed Chinese OCR":
                continue
            key = (exclusion["url"], exclusion.get("member", ""))
            scan_groups.setdefault(key, set()).add(exclusion["page"])
        scans = [{"source_id": self.source_id, "source_url": url, "archive_member": member, "pages": sorted(pages)}
                 for (url, member), pages in sorted(scan_groups.items())]
        summary["image_only_pdf_documents"] = len(scans)
        summary["image_only_pdf_pages"] = sum(len(item["pages"]) for item in scans)
        (output / "scan_inventory.json").write_text(json.dumps(scans, ensure_ascii=False, indent=2), encoding="utf-8")
        with (output / "scan_inventory.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["source_id", "source_url", "archive_member", "pages"])
            writer.writeheader()
            writer.writerows({**item, "pages": ";".join(str(p) for p in item["pages"])} for item in scans)
        (output / "coverage.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        if not records:
            raise ParserError(f"No price-bearing drug rows extracted; inspect {output / 'coverage.json'} for exclusions")
        issues = [self.issue(ErrorType.VALIDATION_ERROR,
                            f"Public extraction: {len(extracted)} price cells; {len(exclusions)} exclusions. "
                            f"Detailed audit: {output / 'coverage.json'}", stage="parsing",
                            extracted_price_cells=summary["extracted_price_cells"], exclusions=len(exclusions))]
        # Publish older descriptions first, so historical files cannot overwrite the newest product metadata.
        records.sort(key=lambda r: (r.data["raw"].get("published_at") or "", r.row_ref))
        return ParsedDataset(self.source_id, snapshot, records,
                             [profile_rows("presentations", COLUMNS, [r.data for r in records], row_count=len(records))],
                             total_rows=len(records), issues=issues, context={"extraction_directory": str(output), "coverage": summary})

    def normalize_public(self, record: ParsedRecord) -> NormalizedRecord:
        row = record.data
        data = row["raw"]
        code = str(row["code"])
        national = bool(re.fullmatch(r"[XZ][A-Z0-9]+", code))
        pack_quantity, pack_unit = chinese_pack(row.get("pack"))
        if pack_unit == "unit" and "粒" in (row.get("pack") or "") and "胶囊" in (row.get("form") or ""):
            pack_unit = "capsule"
        product = SourceProductData(
            product_code=code, code_type="CN_NHSA_DRUG_CODE" if national else "CN_SOURCE_DISCLOSURE",
            product_level=ProductLevel.PRESENTATION, name=row.get("name"),
            ingredient=row.get("name") if national and code.startswith("X") else None, dosage_form=row.get("form"),
            strength_text=row.get("strength"), pack_text=row.get("pack"), pack_quantity=pack_quantity, pack_unit=pack_unit,
            manufacturer=row.get("manufacturer"), authorisation_number=row.get("authorisation"),
            product_type="traditional_chinese_medicine" if code.startswith("Z") else "chemical_or_biological" if national else "drug",
            detail_url=data["notice_url"], notes="Public provincial disclosure, not verified current coverage. "
            f"Original monetary value {data['amount']} ({data['price_label']}, {data.get('value_kind', 'price')}); "
            f"unit: {data.get('pricing_unit') or 'unresolved'}; {data.get('reason') or 'documented price basis'}. "
            f"Attachment: {data['source_url']}. Full amount extraction retained separately.",
        )
        prices = []
        flags = publication_flags(data)
        if flags:
            prices.append(self._flagged_fact(data, flags, pack_quantity, pack_unit))
        elif flags is not None:
            prices.append(PriceFact(
                amount=Decimal(data["amount"]), currency="CNY", price_type=PriceBasis(data["price_type"]),
                price_status=PriceStatus(data["price_status"]), price_unit=data["fact_unit"],
                unit_text=f"{data['price_label']}; {data.get('pricing_unit') or 'pack (explicit header)'}",
                pack_quantity=pack_quantity, pack_unit=pack_unit,
                price_variant=data["price_variant"], coverage_scope=f"province:{self.province}",
                published_at=date.fromisoformat(data["published_at"]) if data.get("published_at") else None,
                valid_from=date.fromisoformat(data["effective_from"]) if data.get("effective_from") else None,
                date_basis="explicit" if data.get("effective_from") else None,
                notes="Published provincial price, not an observed transaction. "
                      f"Source: {data['source_url']}; column: {data['price_label']}.",
            ))
        return NormalizedRecord(record.row_ref, product, prices)

    def _flagged_fact(self, data: dict, flags: list[str], pack_quantity, pack_unit) -> PriceFact:  # noqa: ANN001
        """Source figure published as-is; unknown semantics get separate series instead of guessed ones."""
        variant = data.get("price_variant") or f"unclassified:{data['price_label']}"
        if not data.get("fact_unit"):
            # e.g. guawang_unit_listing would wrongly claim a per-unit basis.
            variant = variant.replace("_unit_", "_").replace("_pack_", "_")
        never_current = NOT_CURRENT_FLAG in flags
        if never_current:
            variant = f"{variant}:{data.get('status') if data.get('status') != 'final' else 'future'}"
        return PriceFact(
            amount=Decimal(str(data["amount"])), currency="CNY",
            price_type=PriceBasis(data["price_type"]) if data.get("price_type") else PriceBasis.OTHER,
            price_status=PriceStatus(data.get("price_status") or "reference"),
            price_unit=data.get("fact_unit") or "unknown",
            unit_text=f"{data['price_label']}; {data.get('pricing_unit') or 'basis not stated'}",
            pack_quantity=pack_quantity, pack_unit=pack_unit,
            price_variant=variant[:128], coverage_scope=f"province:{self.province}",
            published_at=date.fromisoformat(data["published_at"]) if data.get("published_at") else None,
            valid_from=date.fromisoformat(data["effective_from"]) if data.get("effective_from") else None,
            date_basis="explicit" if data.get("effective_from") else None,
            notes=f"Published with warnings ({', '.join(flags)}): {data['reason']}. Original value "
                  f"{data['original_value']} in column {data['price_label']}; source: {data['source_url']}.",
            quality_flags=tuple(flags), never_current=never_current,
        )


def extend_archive_exclusions(exclusions: list | None, member: str, nested: list | None) -> None:
    if exclusions is not None:
        for item in nested or []:
            path = f"{member}/{item['member']}" if item.get("member") else member
            exclusions.append({**item, "member": path})


def price_tables(path: Path, suffix: str, *, exclusions: list | None = None, depth: int = 0):
    """Yield every price-bearing sheet/table, with actual occupied header width."""
    if suffix in {".xls", ".doc"} and zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        if "xl/workbook.xml" in names:
            suffix = ".xlsx"
        elif "word/document.xml" in names:
            suffix = ".docx"
    if suffix in {".rar", ".7z"}:
        try:
            import libarchive
        except OSError as exc:
            raise SchemaChangedError(f"Native libarchive unavailable; archive retained raw: {exc}") from exc
        if depth >= 3:
            raise SchemaChangedError("Archive nesting limit exceeded; retained raw")
        try:
            with libarchive.file_reader(str(path)) as archive, TemporaryDirectory(prefix="cn-prices-") as temp:
                total = 0
                for index, member in enumerate(archive):
                    if index >= 1000:
                        raise SchemaChangedError("Archive member safety limit exceeded; retained raw")
                    if member.isdir:
                        continue
                    part = Path(member.pathname)
                    if not member.isreg or part.is_absolute() or ".." in part.parts:
                        raise SchemaChangedError("Unsafe archive member path/type; retained raw")
                    if member.size and total + member.size > 200_000_000:
                        raise SchemaChangedError("Archive size safety limit exceeded; retained raw")
                    child = Path(temp) / f"{index}{part.suffix.lower()}"
                    with child.open("wb") as handle:
                        for block in member.get_blocks():
                            total += len(block)
                            if total > 200_000_000:
                                raise SchemaChangedError("Archive expanded-size safety limit exceeded; retained raw")
                            handle.write(block)
                    child_exclusions = [] if exclusions is not None else None
                    try:
                        for table in price_tables(child, part.suffix.lower(), exclusions=child_exclusions, depth=depth + 1):
                            yield table._replace(name=f"{member.pathname}/{table.name}")
                    except TABLE_ERRORS as exc:
                        if exclusions is None:
                            raise
                        exclusions.append({"member": member.pathname, "reason": str(exc)})
                    finally:
                        extend_archive_exclusions(exclusions, member.pathname, child_exclusions)
        except libarchive.exception.ArchiveError as exc:
            raise ParserError(f"Archive read failed; retained raw: {exc}") from exc
    elif suffix == ".zip":
        if depth >= 3:
            raise SchemaChangedError("Archive nesting limit exceeded; retained raw")
        with zipfile.ZipFile(path) as archive, TemporaryDirectory(prefix="cn-prices-") as temp:
            members = [m for m in archive.infolist() if not m.is_dir()]
            if len(members) > 1000 or sum(m.file_size for m in members) > 200_000_000:
                raise SchemaChangedError("Archive safety limit exceeded; retained raw")
            for index, member in enumerate(members):
                part = Path(member.filename)
                if part.is_absolute() or ".." in part.parts or (member.external_attr >> 16) & 0o170000 == 0o120000:
                    raise SchemaChangedError("Unsafe archive member path/symlink; retained raw")
                if member.flag_bits & 1:
                    raise SchemaChangedError("Encrypted archive member; retained raw without attempting access")
                child = Path(temp) / f"{index}{part.suffix.lower()}"
                child_exclusions = [] if exclusions is not None else None
                try:
                    child.write_bytes(archive.read(member))
                    for table in price_tables(child, part.suffix.lower(), exclusions=child_exclusions, depth=depth + 1):
                        yield table._replace(name=f"{member.filename}/{table.name}")
                except TABLE_ERRORS as exc:
                    if exclusions is None:
                        raise
                    exclusions.append({"member": member.filename, "reason": str(exc)})
                finally:
                    extend_archive_exclusions(exclusions, member.filename, child_exclusions)
    elif suffix == ".doc":
        converter = shutil.which("textutil")
        if converter is None:
            raise SchemaChangedError("Legacy DOC requires textutil; converter unavailable, retained raw")
        try:
            result = subprocess.run([converter, "-convert", "html", "-stdout", str(path.resolve())],
                                    capture_output=True, timeout=60, check=False)
        except subprocess.TimeoutExpired as exc:
            raise ParserError("Legacy DOC conversion timed out; retained raw") from exc
        if result.returncode:
            raise ParserError(f"Legacy DOC conversion failed: {result.stderr.decode('utf-8', errors='replace')[:300]}")
        soup = BeautifulSoup(result.stdout, "html.parser")
        for index, table in enumerate(soup.select("table")):
            rows = [[compact(cell.get_text(" ", strip=True)) for cell in row.select("td, th")] for row in table.select("tr")]
            for start, header in enumerate(rows[:30]):
                if any(is_price_column(t) for t in header) and any(t in {compact(a) for a in ALIASES["name"]} for t in header):
                    yield PriceTable(f"table{index}", header, list(enumerate(rows[start + 1:], start + 2)),
                                     " ".join(" ".join(row) for row in rows[:start]))
                    break
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > 20_000_000:
                raise SchemaChangedError("DOCX XML safety limit exceeded; retained raw")
            root = ElementTree.fromstring(archive.read(info))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for index, table in enumerate(root.findall(".//w:tbl", ns)):
            rows = [[compact("".join(c.itertext())) for c in r.findall("w:tc", ns)] for r in table.findall("w:tr", ns)]
            for start, header in enumerate(rows[:30]):
                if any(is_price_column(t) for t in header) and any(t in {compact(a) for a in ALIASES["name"]} for t in header):
                    yield PriceTable(f"table{index}", header, list(enumerate(rows[start + 1:], start + 2)),
                                     " ".join(" ".join(row) for row in rows[:start]))
                    break
    elif suffix == ".xlsx":
        with path.open("rb") as handle, closing(openpyxl.load_workbook(handle, read_only=True, data_only=True)) as book:
            for sheet in book:
                header = None
                header_number = None
                caption = []
                # Read full-width headers once, then bound data traversal to actual occupied columns.
                for number, cells in enumerate(sheet.iter_rows(max_row=30, values_only=True), 1):
                    vals = list(cells)
                    while vals and vals[-1] is None:
                        vals.pop()
                    if not vals:
                        continue
                    texts = [compact(v) for v in vals]
                    if any(is_price_column(t) for t in texts) and any(t in {compact(a) for a in ALIASES["name"]} for t in texts):
                        header, header_number = texts, number
                        break
                    caption.extend(t for t in texts if t)
                if header is not None:
                    buffered = [(number, list(cells)) for number, cells in enumerate(
                        sheet.iter_rows(min_row=header_number + 1, max_col=len(header), values_only=True),
                        header_number + 1,
                    ) if any(c is not None for c in cells)]
                    yield PriceTable(sheet.title, header, buffered, " ".join(caption))
    elif suffix == ".xls":
        import xlrd

        with xlrd.open_workbook(str(path), on_demand=True) as book:
            for sheet in book.sheets():
                for index in range(min(30, sheet.nrows)):
                    header = [compact(v) for v in sheet.row_values(index)]
                    if any(is_price_column(t) for t in header) and any(t in {compact(a) for a in ALIASES["name"]} for t in header):
                        caption = " ".join(compact(v) for i in range(index) for v in sheet.row_values(i) if v)
                        yield PriceTable(sheet.name, header, [(i + 1, sheet.row_values(i)) for i in range(index + 1, sheet.nrows)], caption)
                        break
    elif suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for number, page in enumerate(pdf.pages, 1):
                text = compact(page.extract_text() or "")
                if not text and exclusions is not None:
                    exclusions.append({"page": number, "reason": "Image-only PDF page; requires reviewed Chinese OCR"})
                if not PRICE_RE.search(text) or not any(compact(alias) in text for alias in ALIASES["name"]):
                    continue
                for index, table in enumerate(page.extract_tables()):
                    for start, row in enumerate(table[:5]):
                        header = [compact(v) for v in row]
                        if any(is_price_column(t) for t in header) and any(t in {compact(a) for a in ALIASES["name"]} for t in header):
                            yield PriceTable(f"p{number}t{index}", header, [(i + 1, r) for i, r in enumerate(table[start + 1:], start + 1)],
                                             " ".join(compact(v) for row in table[:start] for v in row if v))
                            break
    elif suffix == ".csv":
        import csv

        from app.core.parsing import decode_bytes

        text, _ = decode_bytes(path.read_bytes())
        rows = list(csv.reader(text.splitlines()))
        if rows:
            yield PriceTable("csv", [compact(v) for v in rows[0]], list(enumerate(rows[1:], 2)))
    else:
        raise SchemaChangedError(f"Unsupported attachment format {suffix}; retained raw, not silently discarded")


def extract_table(header: list[str], rows: list, meta: dict, prefix: str) -> tuple[list[ParsedRecord], list[dict]]:
    if non_drug_table(header, meta):
        return [], []
    positions = {}
    for field, aliases in ALIASES.items():
        matched = [i for i, text in enumerate(header) if text in {compact(a) for a in aliases}]
        if matched:
            positions[field] = matched[0]
    if "name" not in positions:
        return [], []
    drug_schema = "药" in header[positions["name"]] or header[positions["name"]] in {"通用名", "通用名称"} or "form" in positions
    drug_context = "药" in (meta.get("title", "") + meta.get("label", ""))
    price_columns = [(i, label) for i, label in enumerate(header) if is_price_column(label)]
    records, prices = [], []
    for number, cells in rows:
        def field(name: str) -> str | None:
            index = positions.get(name)
            return clean_str(cells[index]) if index is not None and index < len(cells) else None

        name = field("name")
        if not name or compact(name) in {compact(a) for a in ALIASES["name"]}:
            continue
        national = compact(field("code"))
        if not drug_schema and not drug_context and not re.fullmatch(r"[XZ][A-Z0-9]+", national) and not re.search(r"国药准字|药注册", field("authorisation") or ""):
            continue
        pack = field("pack")
        if not pack and field("quantity") and field("unit") and field("pack_unit"):
            quantity, quantity_issue = original_amount(field("quantity"))
            if not quantity_issue:
                pack = f"{quantity}{field('unit')}/{field('pack_unit')}"
        raw = {f"{i + 1}:{label}": cells[i] if i < len(cells) else None for i, label in enumerate(header)}
        for index, label in price_columns:
            if index >= len(cells) or clean_str(cells[index]) is None:
                continue
            value = cells[index]
            amount, amount_issue = original_amount(value)
            # Multiple price columns (including historical/proposed) are retained independently.
            previous = bool(re.search(r"调整前|变更前|降价前|调价前|原(?:挂网|中选|采购|入市|价|包)|旧价|历史", label))
            proposed = meta["status"] == "proposed" or bool(re.search(r"拟|申报|报价", label))
            caption = meta.get("table_caption", "")
            withdrawn = meta["status"] == "withdrawn" or bool(re.search(r"废止|作废|撤网", caption))
            template = bool(re.search(r"模板|模版|样表|示例|范例|操作手册|操作指引|用户手册|使用手册|操作指南|演示",
                                      meta.get("label", "") + caption))
            future = any(meta.get(k) and date.fromisoformat(meta[k]) > date.today() for k in ("effective_from", "published_at"))
            semantics = price_semantics(label, field("pricing_unit"), caption, field("unit"))
            explicit_basis = bool(semantics["currency"] and semantics["fact_unit"] and semantics["price_type"])
            ambiguous = header.count(label) > 1
            raw_herb = bool(re.search(r"饮片|原料药|药材", meta.get("title", "")))
            reason = amount_issue
            if reason is None:
                if template:
                    reason = "template/example value"
                elif ambiguous:
                    reason = "ambiguous duplicate price column"
                elif raw_herb:
                    reason = "raw herbal/ingredient material"
                elif semantics["value_kind"] == "aggregate_amount":
                    reason = "Aggregate monetary amount, not a unit price"
                elif semantics["value_kind"] == "daily_cost":
                    reason = "Daily treatment cost or mixed cost/bid-limit basis, not a unit price"
                elif previous or proposed or withdrawn or future:
                    reason = NOT_CURRENT_REASON
                elif not explicit_basis:
                    reason = UNESTABLISHED_REASON
            ref = f"{prefix}:r{number}:c{index + 1}"
            data = {
                "amount": amount, "original_value": value, "price_label": label, "pricing_unit": field("pricing_unit"),
                **semantics, "table_caption": caption,
                "status": "template" if template else "previous" if previous else "proposed" if proposed else "withdrawn" if withdrawn else "final",
                "published_at": meta.get("published_at"), "effective_from": meta.get("effective_from") if not previous else None,
                "source_url": meta["url"], "notice_url": meta["notice_url"], "row_ref": ref,
                "raw": raw, "normalized": reason is None, "reason": reason,
            }
            prices.append(data)
            source_identity = urlparse(meta["url"])._replace(fragment="")
            if "/attachement/" in source_identity.path:
                source_identity = source_identity._replace(query="")
            table_identity = prefix.partition(":")[2] or prefix
            identity = national if national and re.fullmatch(r"[XZ][A-Z0-9]+", national) else "DISCLOSURE:" + hashlib.sha256(
                f"{source_identity.geturl()}|{table_identity}|{number}".encode()
            ).hexdigest()[:32]
            data["product"] = {
                "code": identity, "name": name, "form": field("form"), "strength": field("strength"),
                "pack": pack, "manufacturer": field("manufacturer"), "authorisation": field("authorisation"),
            }
            # One source product can have multiple original price columns, each retained in raw extraction.
            records.append(ParsedRecord(ref, identity, {
                **data["product"], "amount": amount, "raw": data,
            }, role="presentations"))
    return records, prices
