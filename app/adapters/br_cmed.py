"""Brazil – ANVISA/CMED price lists (PMC and PMVG workbooks).

Workbooks have a 'Cabeçalho' sheet (validity date) and a data sheet whose header row is located dynamically.
PMVG values may be stored as Excel formulas (=ROUND(PF*(1-CAP),2)): cached values are used when present,
otherwise the formula is evaluated and the price is flagged price_status=calculated.
Licence CC BY-ND 3.0: internal processing only.
"""

from __future__ import annotations

import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from app.adapters._common import join_notes, pack_price, parse_fields
from app.core.adapter import SourceAdapter
from app.core.enums import PriceBasis, PriceStatus, ProductLevel
from app.core.errors import DiscoveryError, ParserError
from app.core.parsing import clean_str, dedupe_headers, html_links, month_number, parse_decimal, profile_rows, slug
from app.core.schema import ExpectedSchema
from app.core.types import ArtifactFile, NormalizedRecord, ParsedDataset, ParsedRecord, RawSnapshot, SourceArtifact, SourceProductData
from app.normalization.forms import compile_extra
from app.observability import Progress

PAGE = "https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos"
FILE_RE = re.compile(r"lista_(?P<kind>pmc|pmvg)_{1,2}(?P<date>\d{8})_\d+\.xlsx", re.IGNORECASE)
VALID_RE = re.compile(r"v[áa]lidos a partir de (\d{1,2}) de ([a-zç]+) de (\d{4})", re.IGNORECASE)
FORMULA_RE = re.compile(r"^=ROUND\(\s*([A-Z]{1,3})(\d+)\s*\*\s*\(\s*1\s*-\s*([\d.]+)\s*\)\s*,\s*2\s*\)$", re.IGNORECASE)
REQUIRED = ["substancia", "laboratorio", "codigo_ggrem", "registro", "ean_1", "produto", "apresentacao", "classe_terapeutica", "regime_de_preco", "pf_sem_impostos"]
BR_FORMS = compile_extra(
    [
        ("prolonged-release tablet", [r"com rev lib prol", r"com lib prol", r"com lib ret"]),
        ("gastro-resistant tablet", [r"com rev lib retard", r"com rev gastro"]),
        ("effervescent tablet", [r"com efev"]),
        ("orodispersible tablet", [r"com orodisp"]),
        ("tablet", [r"com", r"com rev", r"com sulc", r"comp"]),
        ("capsule", [r"cap dura", r"cap mole", r"cap gel", r"cap"]),
        ("powder for injection", [r"po liof sol inj", r"po sol inj", r"po liof"]),
        ("injection", [r"sol inj", r"sus inj", r"emu inj", r"sol inf", r"inj"]),
        ("eye drops", [r"sol oft", r"sus oft", r"col"]),
        ("nasal spray", [r"spr nas", r"sol nas"]),
        ("inhalation", [r"aer", r"po inal", r"sol inal", r"sus aer"]),
        ("oral solution", [r"sol or", r"sol oral"]),
        ("oral suspension", [r"sus or", r"po sus or"]),
        ("syrup", [r"xpe"]),
        ("drops", [r"sol or gts", r"gts"]),
        ("cream", [r"crem derm", r"crem vag", r"crem"]),
        ("ointment", [r"pom derm", r"pom oft", r"pom"]),
        ("gel", [r"gel"]),
        ("patch", [r"ades", r"adesivo"]),
        ("suppository", [r"sup"]),
        ("granules", [r"gran", r"gran efev"]),
        ("powder", [r"po"]),
    ]
)


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


class BrCmedAdapter(SourceAdapter):
    adapter_id = "br_cmed"
    parser_type = "xlsx"
    parser_version = "1.0"

    def discover(self) -> SourceArtifact:
        page = self.ctx.http.get_text(PAGE)
        latest: dict[str, tuple[str, str]] = {}
        for url, _ in html_links(page.text, PAGE):
            m = FILE_RE.search(url)
            if m:
                kind = m.group("kind").lower()
                if kind not in latest or m.group("date") > latest[kind][0]:
                    latest[kind] = (m.group("date"), url)
        if "pmc" not in latest:
            raise DiscoveryError("CMED page: no 'lista_pmc_<date>.xlsx' link found", details={"page": PAGE})
        files = [
            ArtifactFile(url=url, filename=f"lista_{kind}_{d}.xlsx", role=kind)
            for kind, (d, url) in sorted(latest.items())
        ]
        d = max(v[0] for v in latest.values())
        published = date(int(d[:4]), int(d[4:6]), int(d[6:]))
        notes = [] if "pmvg" in latest else ["PMVG list not found on the page – only PMC/PF prices will be loaded"]
        return SourceArtifact(self.source_id, files, published_at=published, edition=f"Lista CMED {published:%d/%m/%Y}", discovered_from=PAGE, notes=notes)

    def expected_schemas(self) -> list[ExpectedSchema]:
        return [
            ExpectedSchema("pmc", REQUIRED + ["pmc_sem_impostos"], typed_columns=["codigo_ggrem"]),
            ExpectedSchema("pmvg", REQUIRED + ["pmvg_sem_impostos"], typed_columns=["codigo_ggrem"]),
        ]

    @staticmethod
    def _read_workbook(path: Path, stop: int | None) -> tuple[list[str], list[dict], date | None, int, bool]:
        import openpyxl

        values_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas_wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            valid_from = None
            for name in values_wb.sheetnames:
                ws = values_wb[name]
                for row in ws.iter_rows(max_row=10, values_only=True):
                    for cell in row:
                        m = VALID_RE.search(str(cell or ""))
                        if m and month_number(m.group(2)):
                            valid_from = date(int(m.group(3)), month_number(m.group(2)), int(m.group(1)))
            data_sheet = None
            header_idx = None
            for name in values_wb.sheetnames:
                for i, row in enumerate(values_wb[name].iter_rows(max_row=60, values_only=True)):
                    if row and "substancia" in [slug(c) for c in row if c] and "codigo_ggrem" in [slug(c) for c in row if c]:
                        data_sheet, header_idx = name, i
                        break
                if data_sheet:
                    break
            if data_sheet is None:
                raise ParserError(f"{path.name}: no sheet with a 'SUBSTÂNCIA … CÓDIGO GGREM' header row")
            values_rows = values_wb[data_sheet].iter_rows(values_only=True)
            formula_rows = formulas_wb[data_sheet].iter_rows(values_only=True)
            header: list[str] = []
            out: list[dict] = []
            total = 0
            truncated = False
            progress = Progress(f"read {path.name} [{data_sheet}]", unit="rows", check_every=500)
            for i, (vals, forms) in enumerate(zip(values_rows, formula_rows)):
                progress.update()
                if i < header_idx:
                    continue
                if i == header_idx:
                    header = dedupe_headers(vals)
                    continue
                if not any(v not in (None, "") for v in vals):
                    continue
                total += 1
                if stop is not None and len(out) >= stop:
                    truncated = True
                    break
                row = dict(zip(header, vals))
                for j, (v, f) in enumerate(zip(vals, forms)):
                    if v is None and isinstance(f, str) and f.startswith("="):
                        m = FORMULA_RE.match(f.replace(" ", ""))
                        if not m:
                            row[header[j]] = f  # left as-is: parse_decimal will raise an explicit error
                            continue
                        ref = vals[_col_index(m.group(1))]
                        base = parse_decimal(ref, decimal_sep=",")
                        if base is not None:
                            calc = (base * (1 - Decimal(m.group(3)))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                            row[header[j]] = calc
                            row.setdefault("_calculated", []).append(header[j])
                            row["_cap"] = m.group(3)
                row["_line"] = i + 1
                out.append(row)
            progress.finish(kept=len(out), formulas_evaluated=sum(1 for r in out if r.get("_calculated")))
            return header, out, valid_from, total, truncated
        finally:
            values_wb.close()
            formulas_wb.close()

    def parse(self, snapshot: RawSnapshot) -> ParsedDataset:
        stop = self.hints.early_stop_at
        header, pmc_rows, valid_from, total, truncated = self._read_workbook(snapshot.file("pmc").path, stop)
        profiles = [profile_rows("pmc", header, pmc_rows, row_count=total)]
        pmvg_by_ggrem: dict[str, dict] = {}
        try:
            pmvg_file = snapshot.file("pmvg")
        except KeyError:
            pmvg_file = None
        if pmvg_file is not None:
            h2, pmvg_rows, _, total2, trunc2 = self._read_workbook(pmvg_file.path, stop * 3 if stop else None)
            profiles.append(profile_rows("pmvg", h2, pmvg_rows, row_count=total2))
            truncated = truncated or trunc2
            pmvg_by_ggrem = {str(r.get("codigo_ggrem")).strip(): r for r in pmvg_rows if r.get("codigo_ggrem")}
        records = []
        for r in pmc_rows:
            key = clean_str(r.get("codigo_ggrem"))
            data = {**r, "_pmvg": pmvg_by_ggrem.get(key or "")}
            records.append(ParsedRecord(row_ref=f"{snapshot.file('pmc').filename}:{r['_line']}", key=key, data=data))
        return ParsedDataset(self.source_id, snapshot, records, profiles, truncated=truncated, total_rows=total, context={"valid_from": valid_from})

    def record_matches_product_ids(self, record: ParsedRecord, product_ids: set[str]) -> bool:
        return record.key in product_ids or clean_str(record.data.get("ean_1")) in product_ids

    def normalize_record(self, record: ParsedRecord, dataset: ParsedDataset) -> NormalizedRecord | None:
        d = record.data
        apres = clean_str(d.get("apresentacao")) or ""
        regime = clean_str(d.get("regime_de_preco"))
        status = PriceStatus.CEILING if (regime or "").lower().startswith("regulado") else PriceStatus.REFERENCE
        valid_from = dataset.context.get("valid_from") or dataset.snapshot.published_at
        hospital = clean_str(d.get("restricao_hospitalar"))
        product = SourceProductData(
            product_code=clean_str(d.get("codigo_ggrem")) or "",
            code_type="GGREM",
            product_level=ProductLevel.PRESENTATION,
            name=clean_str(f"{d.get('produto') or ''} {apres}"),
            brand_name=clean_str(d.get("produto")),
            ingredient=clean_str(d.get("substancia")),
            strength_text=apres or None,
            pack_text=apres or None,
            manufacturer=clean_str(d.get("laboratorio")),
            therapeutic_class=clean_str(d.get("classe_terapeutica")),
            product_type=clean_str(d.get("tipo_de_produto_status_do_produto")),
            authorisation_number=clean_str(d.get("registro")),
            ean=", ".join(e for e in (clean_str(d.get(f"ean_{i}")) for i in (1, 2, 3)) if e and e != "-") or None,
            legal_status=clean_str(d.get("tarja")),
            restrictions=join_notes(
                f"regime de preço: {regime}" if regime else None,
                "restrição hospitalar" if (hospital or "").lower() in ("sim", "s") else None,
            ),
            notes=f"CNPJ {d['cnpj']}" if clean_str(d.get("cnpj")) else None,
            **parse_fields(strength_text=apres, form_texts=[apres], pack_text=apres, extra_forms=BR_FORMS),
        )
        common = dict(currency="BRL", product=product, valid_from=valid_from, coverage_scope=self.config.coverage_scope)
        prices = [
            pack_price(parse_decimal(d.get("pf_sem_impostos"), decimal_sep=","), price_type=PriceBasis.EX_FACTORY, status=status,
                       vat_included=False, variant="pf_sem_impostos", notes="sem impostos", **common),
        ]
        for rate in self.option("icms_rates", ["18 %"]):
            col = slug(rate)
            tax = parse_decimal(rate.replace("%", ""), decimal_sep=",")
            prices.append(pack_price(parse_decimal(d.get(f"pf_{col}"), decimal_sep=","), price_type=PriceBasis.EX_FACTORY, status=status,
                                     vat_included=True, variant=f"pf_icms_{col}", tax_rate=tax, **common))
            prices.append(pack_price(parse_decimal(d.get(f"pmc_{col}"), decimal_sep=","), price_type=PriceBasis.RETAIL,
                                     status=PriceStatus.CEILING, vat_included=True, variant=f"pmc_icms_{col}", tax_rate=tax, **common))
        pmvg = d.get("_pmvg") or {}
        calculated = set(pmvg.get("_calculated") or [])
        for col, variant, vat, rate in [("pmvg_sem_impostos", "pmvg_sem_impostos", False, None)] + [
            (f"pmvg_{slug(r)}", f"pmvg_icms_{slug(r)}", True, r) for r in self.option("icms_rates", ["18 %"])
        ]:
            amount = parse_decimal(pmvg.get(col), decimal_sep=",")
            note = f"PF × (1 − CAP {pmvg.get('_cap')}) evaluated from workbook formula" if col in calculated else None
            prices.append(pack_price(amount, price_type=PriceBasis.PUBLIC_PROCUREMENT,
                                     status=PriceStatus.CALCULATED if col in calculated else PriceStatus.CEILING,
                                     vat_included=vat, variant=variant, notes=note,
                                     tax_rate=parse_decimal(rate.replace("%", ""), decimal_sep=",") if rate else None, **common))
        return NormalizedRecord(record.row_ref, product, [p for p in prices if p is not None])
