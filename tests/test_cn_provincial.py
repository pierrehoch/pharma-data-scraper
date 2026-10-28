"""Synthetic, independently authored tables; no redistribution of restricted workbooks."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import func, select

from app.adapters._cn_provincial import GD_COLUMNS, HB_ADMISSION_COLUMNS, HB_ADJUSTMENT_COLUMNS, HB_AWARD_COLUMNS, chinese_pack
from app.core.enums import JobStatus, PriceBasis, PriceStatus
from app.core.errors import DiscoveryError, ParserError, SchemaChangedError, SourceUnsupportedError
from app.core.pipeline import RunOptions, Selection
from app.core.pricing import normalize_price
from app.db.models import Price, Product
from tests.test_adapters import make_adapter

CODE = "XA01ABC000A001010100001"


def workbook(columns, rows, *, inflated=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Attachment"])
    sheet.append(["Synthetic test disclosure"])
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(c) for c in columns])
    if inflated:
        sheet.cell(8, 16383).number_format = "0.00"
    book.create_sheet("empty")
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def gd_row(**changes):
    return {
        "序号": 1, "批次": "test", "药品ID": 1001, "国家医保代码": CODE, "产品名称": "测试片",
        "剂型名称": "片剂", "具体规格": "10mg", "包装规格": "28片/盒", "包装材质": "test",
        "生产企业/厂商": "Test manufacturer", "申报企业名称": "Different declaring company",
        "批准文号/注册证号": "TEST-H001", "包装挂网价格（元）": 57.76, **changes,
    }


def hb_row(**changes):
    return {
        "序号": 1, "流水号": 2001, "国家医保编码": CODE, "药品名称": "测试片", "商品名称": "无",
        "剂型": "片剂", "规格": "10mg", "最小包装数量": "28", "最小包装单位": "盒",
        "最小制剂单位": "片", "包装材质": "test", "批准文号": "TEST-H001",
        "上市许可持有人": "Test holder", "生产企业名称": "Test manufacturer",
        "统一社会信用代码": "test", "申报企业名称": "Different declaring company",
        "申报类型": "普通挂网药品-其他药品", "质量层次": "test", "价格形成方式": 5,
        "备注": None, "挂网价": 58.76, **changes,
    }


def configure(adapter, http, columns=None, rows=None, **changes):
    adapter.config.options["internal_research_acknowledged"] = True
    doc = dict(adapter.default_documents[0], **changes)
    adapter.config.options["documents"] = [doc]
    if columns is not None:
        http.override(r"\.xlsx$", workbook(columns, rows or []))
    return doc


def ingest(adapter):
    snapshot = adapter.download(adapter.discover())
    parsed = adapter.parse(snapshot)
    normalized = adapter.normalize(parsed)
    assert normalized.failed_rows == 0
    assert not adapter.validate(normalized).errors
    return snapshot, parsed, normalized


@pytest.mark.parametrize("source_id", ["cn_gdmede", "cn_hubei"])
def test_opt_in_and_disabled_defaults(source_id, settings, http, registry):
    adapter = make_adapter(source_id, settings, http, registry)
    assert not adapter.config.enabled
    adapter.config.options["internal_research_acknowledged"] = False
    assert adapter.check_configuration()
    with pytest.raises(SourceUnsupportedError, match="internal_research_acknowledged"):
        adapter.discover()
    assert not http.requested


def test_guangdong_exact_pack_price_and_snapshot_interpretation(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    doc = configure(adapter, http, GD_COLUMNS, [gd_row()])
    snapshot = adapter.download(adapter.discover())
    # A config edit must not change the interpretation of an already saved snapshot.
    adapter.config.options["documents"][0]["published_at"] = "2020-01-01"
    normalized = adapter.normalize(adapter.parse(snapshot))
    product = normalized.records[0].product
    fact = normalized.records[0].prices[0]
    assert product.product_code == CODE and product.manufacturer == "Test manufacturer"
    assert product.authorisation_number == "TEST-H001"
    assert fact.amount == Decimal("57.76") and fact.currency == "CNY"
    assert fact.price_unit == "pack" and fact.pack_quantity == 28 and fact.pack_unit == "tablet"
    assert fact.price_type == PriceBasis.PUBLIC_PROCUREMENT and fact.price_status == PriceStatus.REFERENCE
    assert fact.price_variant == "guawang_pack_listing" and fact.coverage_scope == "province:CN-GD"
    assert fact.valid_from is None and fact.published_at == date(2026, 10, 13)
    assert normalize_price(fact) == (Decimal("2.06285714"), "tablet")
    assert snapshot.file("reviewed_metadata").path.is_file()
    assert str(doc["url"]) in fact.notes


def test_hubei_withholds_unverified_prices(settings, http, registry):
    adapter = make_adapter("cn_hubei", settings, http, registry)
    configure(adapter, http, HB_ADMISSION_COLUMNS, [hb_row()])
    _, parsed, normalized = ingest(adapter)
    assert not normalized.records[0].prices
    assert any("has not been verified" in i.message for i in parsed.issues)
    assert normalized.records[0].product.pack_quantity == 28
    assert normalized.records[0].product.brand_name is None


def test_hubei_reviewed_price_and_provincial_difference(settings, http, registry):
    gd = make_adapter("cn_gdmede", settings, http, registry)
    configure(gd, http, GD_COLUMNS, [gd_row()])
    _, _, gd_data = ingest(gd)
    hb = make_adapter("cn_hubei", settings, http, registry)
    configure(hb, http, HB_ADMISSION_COLUMNS, [hb_row()],
              cny_pack_evidence_url="https://ybj.hubei.gov.cn/reviewed-unit-rule.shtml")
    _, _, hb_data = ingest(hb)
    assert gd_data.records[0].product.product_code == hb_data.records[0].product.product_code
    gd_price, hb_price = gd_data.records[0].prices[0], hb_data.records[0].prices[0]
    assert gd_price.amount != hb_price.amount
    assert gd_price.coverage_scope != hb_price.coverage_scope
    assert hb_price.valid_from == date(2026, 10, 12)


@pytest.mark.parametrize("changes,reason", [
    ({"status": "proposed", "effective_from": None}, "proposed"),
    ({"effective_from": "2999-01-01"}, "future effective"),
    ({"published_at": "2999-01-01"}, "future publication"),
])
def test_non_current_documents_are_product_only(settings, http, registry, changes, reason):
    adapter = make_adapter("cn_hubei", settings, http, registry)
    configure(adapter, http, HB_ADMISSION_COLUMNS, [hb_row()], **changes)
    _, parsed, normalized = ingest(adapter)
    assert not normalized.records[0].prices
    assert any(reason in i.message for i in parsed.issues)


@pytest.mark.parametrize("pack,expected", [
    ("28片/盒", (Decimal(28), "tablet")), ("1支/盒", (Decimal(1), "unit")),
    ("60粒/瓶", (Decimal(60), "unit")), ("2×14片/盒", (None, None)),
    ("1盒", (None, None)), ("0片/盒", (None, None)),
])
def test_chinese_pack_does_not_guess(pack, expected):
    assert chinese_pack(pack) == expected


def test_traditional_specialties_not_mapped_to_single_ingredients(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(**{"国家医保代码": "ZA09TEST001", "产品名称": "测试丸", "包装规格": "380丸/盒"})])
    _, _, normalized = ingest(adapter)
    assert normalized.records[0].product.ingredient is None
    assert normalized.records[0].product.product_type == "traditional_chinese_medicine"
    assert normalized.records[0].prices[0].pack_quantity == 380


@pytest.mark.parametrize("columns", [
    GD_COLUMNS[:-1],
    ["序号", "流水号", "药品代码", "通用名", "剂型", "规格", "包装数量", "制剂单位",
     "包装材质", "计价单位", "批准文号", "生产企业", "上市许可人", "申报企业"],
    ["序号", "流水号", "药品代码", "药品名称", "剂型", "包装规格", "包装材质",
     "包装数量", "制剂单位", "生产企业", "中选企业", "包装单位", "备注"],
])
def test_supplier_only_or_proposed_lists_never_create_prices(settings, http, registry, columns):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, columns, [gd_row()])
    with pytest.raises(SchemaChangedError, match="headers"):
        ingest(adapter)


def test_inflated_dimensions_do_not_expand_empty_columns(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http)
    http.override(r"\.xlsx$", workbook(GD_COLUMNS, [gd_row()], inflated=True))
    _, parsed, normalized = ingest(adapter)
    assert parsed.total_rows == 1 and len(normalized.records) == 1


@pytest.mark.parametrize("value", [None, 0, -1, True, "not a price"])
def test_invalid_amount_is_explicit_error(settings, http, registry, value):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(**{"包装挂网价格（元）": value})])
    with pytest.raises(ParserError, match="price"):
        ingest(adapter)


@pytest.mark.parametrize("changes", [
    {"url": "https://example.com/data.xlsx"},
    {"notice_url": "http://www.gdmede.com.cn/notice"},
    {"url": "https://www.gdmede.com.cn/data.pdf"},
    {"published_at": "bad date"}, {"status": "unknown"}, {"kind": "hb_admission"},
    {"status": "proposed", "effective_from": "2026-01-01"},
    {"cny_pack_evidence_url": "https://example.com/claim"},
])
def test_document_configuration_rejected(settings, http, registry, changes):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, **changes)
    with pytest.raises(DiscoveryError):
        adapter.discover()


def test_conflicting_product_prices_are_not_silently_overwritten(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(), gd_row(**{"序号": 2, "包装挂网价格（元）": 58})])
    with pytest.raises(ParserError, match="conflicting"):
        ingest(adapter)


def test_identical_disclosures_keep_provenance_without_false_conflict(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    doc = configure(adapter, http, GD_COLUMNS, [gd_row(), gd_row(**{"序号": 2})])
    adapter.config.options["documents"].append(dict(doc, url="https://www.gdmede.com.cn/attachement/another.xlsx"))
    _, parsed, normalized = ingest(adapter)
    assert parsed.total_rows == 4
    assert len(normalized.records) == 4
    assert len({r.row_ref for r in normalized.records}) == 4


def adjustment_row(**changes):
    return {
        "序号": 1, "流水号": 2001, "医保编码": CODE, "药品编码": CODE,
        "申报企业名称": "Different declaring company", "通用名": "测试片", "剂型": "片剂", "规格": "10mg",
        "最小包装数量": "28", "包装数量": 28, "批准文号": "TEST-H001",
        "挂网类型": "省级（际）集中带量采购药品", "上报类型": "价格下调",
        "降价前价格": 58.76, "降价后价格": 50, "备注": None, **changes,
    }


def test_adjustment_enriched_only_by_exact_code_and_no_old_price_dates(settings, http, registry):
    adapter = make_adapter("cn_hubei", settings, http, registry)
    doc = configure(adapter, http, cny_pack_evidence_url="https://ybj.hubei.gov.cn/reviewed-unit-rule.shtml")
    second = dict(adapter.default_documents[1], cny_pack_evidence_url=doc["cny_pack_evidence_url"])
    adapter.config.options["documents"].append(second)
    http.override(re_escape(doc["url"]), workbook(HB_ADMISSION_COLUMNS, [hb_row()]))
    http.override(re_escape(second["url"]), workbook(HB_ADJUSTMENT_COLUMNS, [adjustment_row()]))
    _, _, normalized = ingest(adapter)
    adjustment = normalized.records[1]
    assert adjustment.product.manufacturer == "Test manufacturer"
    assert adjustment.product.pack_text == "28片/盒" and adjustment.product.pack_unit == "tablet"
    assert len(adjustment.prices) == 1 and adjustment.prices[0].amount == 50
    assert adjustment.prices[0].valid_from == date(2026, 10, 13)
    assert "start date unknown" in adjustment.prices[0].notes


def re_escape(text):
    import re
    return re.escape(text)


def test_award_ambiguous_date_is_not_guessed(settings, http, registry):
    adapter = make_adapter("cn_hubei", settings, http, registry)
    configure(adapter, http, HB_AWARD_COLUMNS, [adjustment_row()],
              kind="hb_award_adjustment", effective_from=None,
              cny_pack_evidence_url="https://ybj.hubei.gov.cn/reviewed-unit-rule.shtml")
    _, parsed, normalized = ingest(adapter)
    assert not normalized.records[0].prices
    assert any("ambiguous" in i.message for i in parsed.issues)
    # Declaring enterprise is not a manufacturer.
    assert normalized.records[0].product.manufacturer is None
    assert normalized.records[0].product.pack_unit is None


def test_pipeline_idempotency_history_and_incremental_absence(pipeline, session_factory, http):
    adapter = make_adapter("cn_gdmede", pipeline.settings, http, pipeline.registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(), gd_row(**{"序号": 2, "国家医保代码": "XA01OTHER001"})])
    source = Selection(source_ids=["cn_gdmede"])
    first = pipeline.run(source, RunOptions())
    assert first.jobs[0].status == JobStatus.PARTIAL_SUCCESS
    pipeline.run(source, RunOptions(force_download=True))
    with session_factory() as session:
        assert session.scalar(select(func.count()).select_from(Price)) == 2
    adapter.config.options["documents"][0]["published_at"] = "2026-10-24"
    http.override(r"\.xlsx$", workbook(GD_COLUMNS, [gd_row(**{"包装挂网价格（元）": 60})]))
    pipeline.run(source, RunOptions(accept_schema_change=True))
    with session_factory() as session:
        assert session.scalar(select(func.count()).select_from(Product)) == 2
        prices = list(session.scalars(select(Price).order_by(Price.id)))
        assert len(prices) == 3
        assert sum(p.is_current for p in prices) == 2
        assert prices[0].is_current is False and prices[-1].amount == 60


def test_pipeline_hubei_unit_warning_is_persisted(pipeline, session_factory, http):
    adapter = make_adapter("cn_hubei", pipeline.settings, http, pipeline.registry)
    configure(adapter, http, HB_ADMISSION_COLUMNS, [hb_row()])
    outcome = pipeline.run(Selection(source_ids=["cn_hubei"]), RunOptions()).jobs[0]
    assert outcome.status == JobStatus.PARTIAL_SUCCESS
    with session_factory() as session:
        assert session.scalar(select(func.count()).select_from(Product)) == 1
        assert session.scalar(select(func.count()).select_from(Price)) == 0


def test_api_dashboard_and_csv_keep_provinces_and_limitations(pipeline, http):
    from fastapi.testclient import TestClient

    from app.api.main import create_app

    gd = make_adapter("cn_gdmede", pipeline.settings, http, pipeline.registry)
    gd_doc = configure(gd, http)
    hb = make_adapter("cn_hubei", pipeline.settings, http, pipeline.registry)
    hb_doc = configure(hb, http)
    http.override(re_escape(gd_doc["url"]), workbook(GD_COLUMNS, [gd_row()]))
    http.override(re_escape(hb_doc["url"]), workbook(HB_ADMISSION_COLUMNS, [hb_row()]))
    pipeline.run(Selection(source_ids=["cn_gdmede", "cn_hubei"]), RunOptions())
    with TestClient(create_app()) as client:
        sources = {r["source_id"]: r for r in client.get("/sources?country=CN").json()}
        assert sources["cn_gdmede"]["coverage_scope"] == "province:CN-GD"
        assert sources["cn_hubei"]["prices"] == 0
        assert "unverified" in sources["cn_hubei"]["limitations"]
        rows = client.get("/prices?country=CN").json()["items"]
        gd_row_out = next(r for r in rows if r["source_id"] == "cn_gdmede")
        hb_row_out = next(r for r in rows if r["source_id"] == "cn_hubei")
        assert gd_row_out["amount"] == "57.76" and gd_row_out["coverage_scope"] == "province:CN-GD"
        assert hb_row_out["price_id"] is None and hb_row_out["amount"] is None
        csv = client.get("/prices.csv?source_id=cn_gdmede")
        assert "province:CN-GD" in csv.text and "guawang_pack_listing" in csv.text
        page = client.get("/dashboard/sources/cn_hubei")
        assert page.status_code == 200 and "unverified" in page.text and "province:CN-HB" in page.text


def test_capsule_unit_requires_form_evidence(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(**{"剂型名称": "胶囊剂", "包装规格": "60粒/瓶"})])
    _, _, normalized = ingest(adapter)
    assert normalized.records[0].product.pack_unit == "capsule"


def test_unpublished_strength_is_left_empty_not_invented(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http, GD_COLUMNS, [gd_row(**{"具体规格": "-"})])
    _, _, normalized = ingest(adapter)
    assert normalized.records[0].product.strength_text is None
    assert normalized.records[0].prices[0].amount == Decimal("57.76")


def test_later_backfill_does_not_replace_newer_document_price(pipeline, session_factory, http):
    adapter = make_adapter("cn_gdmede", pipeline.settings, http, pipeline.registry)
    configure(adapter, http, GD_COLUMNS, [gd_row()], published_at="2026-10-24")
    source = Selection(source_ids=["cn_gdmede"])
    pipeline.run(source, RunOptions())
    adapter.config.options["documents"][0]["published_at"] = "2026-10-13"
    http.override(r"\.xlsx$", workbook(GD_COLUMNS, [gd_row(**{"包装挂网价格（元）": 40})]))
    pipeline.run(source, RunOptions())
    with session_factory() as session:
        prices = list(session.scalars(select(Price).order_by(Price.id)))
        assert len(prices) == 2
        assert prices[0].is_current and not prices[1].is_current


def test_proposed_title_cannot_be_labelled_final(settings, http, registry):
    adapter = make_adapter("cn_gdmede", settings, http, registry)
    configure(adapter, http)
    book = openpyxl.load_workbook(BytesIO(workbook(GD_COLUMNS, [gd_row()])))
    book.active["A2"] = "药品拟挂网公示"
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    http.override(r"\.xlsx$", buffer.getvalue())
    with pytest.raises(SchemaChangedError, match="contradicts"):
        ingest(adapter)
